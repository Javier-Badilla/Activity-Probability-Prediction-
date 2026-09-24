import datetime
import io
import math
import re
from dataclasses import dataclass
import pandas as pd
from pathlib import Path
import streamlit as st

# --------------------------------------------------------------------------
# Logo and page
# --------------------------------------------------------------------------


# Findig app.py
BASE_DIR = Path(__file__).parent

LOGO_PATH = BASE_DIR / "assets" / "logo.png"
ICON_PATH = BASE_DIR / "assets" / "icon.ico"

# Page config
st.set_page_config(
    page_title="SPPS Builder | Síntesis en Fase Sólida",
    page_icon=str(ICON_PATH) if ICON_PATH.exists() else "🧪",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Logo
if LOGO_PATH.exists():
    st.logo(
        image=str(LOGO_PATH),
        icon_image=str(ICON_PATH) if ICON_PATH.exists() else None
    )
else:
    st.warning(f"Logo not found at: {LOGO_PATH}")


# --------------------------------------------------------------------------
# Data constants
# --------------------------------------------------------------------------
ONE2THREE = {
    "A": "ALA", "R": "ARG", "N": "ASN", "D": "ASP", "C": "CYS", "Q": "GLN",
    "E": "GLU", "G": "GLY", "H": "HIS", "I": "ILE", "L": "LEU", "K": "LYS",
    "M": "MET", "F": "PHE", "P": "PRO", "S": "SER", "T": "THR", "W": "TRP",
    "Y": "TYR", "V": "VAL",
}

TOTALS_ORDER = ["ALA", "ARG", "ASN", "ASP", "CYS", "PHE", "GLY", "GLU", "GLN", "HIS",
                "ILE", "LEU", "LYS", "MET", "PRO", "SER", "TYR", "THR", "TRP", "VAL"]

RESIDUE_MASS = {
    "A": 71.0788, "R": 156.1875, "N": 114.1038, "D": 115.0886, "C": 103.1388,
    "Q": 128.1307, "E": 129.1155, "G": 57.0519, "H": 137.1411, "I": 113.1594,
    "L": 113.1594, "K": 128.1741, "M": 131.1926, "F": 147.1766, "P": 97.1167,
    "S": 87.0782, "T": 101.1051, "W": 186.2132, "Y": 163.1760, "V": 99.1326,
}
WATER = 18.0153
SPECIAL_MASSES = {
    "orn": 114.1472, "aib": 85.1045, "nle": 113.1594, "nva": 99.1326,
    "dab": 100.1205, "dap": 86.0932, "abu": 85.1045, "cit": 157.1721,
}
ACTIVATORS = {
    "TBTU": ("AA+TBTU+OXYMA+DIPEA", "(5:5:5:7,5)"),
    "HBTU": ("AA+HBTU+OXYMA+DIPEA", "(5:5:5:7,5)"),
    "HCTU": ("AA+HCTU+OXYMA+DIPEA", "(5:5:5:7,5)"),
    "DIC":  ("AA+DIC+OXYMA", "(5:5:5)"),
}
DEFAULT_COUPLINGS = {"simple": "TBTU", "doble": "HBTU", "triple": "HCTU"}
DEFAULT_DEPROTECTION = "PP 20% TritonX100 1%/DMF"
STANDARD = set(ONE2THREE)
PAGEBREAK = "PAGEBREAK"
CHK = "[ ]"


# --------------------------------------------------------------------------
# Lógica Interna y Parser
# --------------------------------------------------------------------------
@dataclass(frozen=True)
class Res:
    key: str
    d: bool = False
    special: bool = False

    @property
    def label(self):
        base = self.key.upper() if self.special else ONE2THREE[self.key]
        return ("D-" if self.d else "") + base

    @property
    def sortkey(self):
        return (self.key.upper(), self.d)

    @property
    def mass(self):
        if self.special:
            return SPECIAL_MASSES.get(self.key.lower())
        return RESIDUE_MASS[self.key]


_TOKEN = re.compile(r"\[([^\]]+)\]|([A-Za-z])")


def parse_sequence(text):
    s = re.sub(r"\s+", "", str(text))
    out, pos = [], 0
    for m in _TOKEN.finditer(s):
        if m.start() != pos:
            raise ValueError(f"Carácter no válido en {s!r}: {s[pos:m.start()]!r}")
        pos = m.end()
        if m.group(1):
            name = m.group(1).strip()
            d = name[:2].lower() == "d-"
            if d:
                name = name[2:]
            out.append(Res(name, d, True))
        else:
            ch = m.group(2)
            if ch.upper() not in STANDARD:
                raise ValueError(f"Aminoácido desconocido {ch!r} en {s!r}")
            out.append(Res(ch.upper(), ch.islower(), False))
    if pos != len(s):
        raise ValueError(f"Carácter no válido en {s!r}: {s[pos:]!r}")
    if not out:
        raise ValueError("Secuencia vacía")
    return out


@dataclass
class Peptide:
    bag: int
    seq: list
    text: str
    family: str = ""
    pos: str = ""

    @property
    def length(self):
        return len(self.seq)

    @property
    def mw(self):
        masses = [r.mass for r in self.seq]
        if any(m is None for m in masses):
            return None
        return sum(masses) + WATER


def _plain(s):
    return [(s, "")]


def build_simultaneous_program(peptides, name, mg, deprotection, couplings, windows_by="posicion"):
    couplings = {**DEFAULT_COUPLINGS, **(couplings or {})}
    peptides = sorted(peptides, key=lambda p: p.bag)
    L = []
    add = lambda s="": L.append(_plain(s))

    families = {p.family for p in peptides if p.family}
    add(f"Nombre de Síntesis: {name}")
    add(f"Cantidad de Síntesis: {mg} mg.")
    add(f"Método de Desprotección: {deprotection} ")
    add(f"Cantidad de Péptidos: {len(peptides)}")
    add(f"Cantidad de Familias: {max(1, len(families))}")
    add("-")
    add("Bolsa   Largo   Pos.    M.W.    Familia")
    for p in peptides:
        mw = f"{p.mw:.2f}" if p.mw is not None else "s/d"
        add(f"{p.bag:<8}{p.length:<8}{p.pos:<8}{mw:<8}{p.family}".rstrip() + " ")
        add(f"Secuencia: {p.text}")
    for p in peptides:
        add(f"{p.bag:<8}{p.text}")

    used = {r for p in peptides for r in p.seq}
    extra = sorted((r for r in used if r.d or r.special), key=lambda r: (r.label))
    rows = [(lab, lambda r, lab=lab: (not r.d and not r.special and r.label == lab))
            for lab in TOTALS_ORDER]
    rows += [(r.label, lambda x, r=r: x == r) for r in extra]

    def table(residues, out_count=None):
        for lab, match in rows:
            n = sum(1 for r in residues if match(r))
            add(f"{lab:<16}{n:<8}RESIDUOS")
        if out_count is not None:
            add(f"{'OUT':<16}{out_count:<8}RESIDUOS")

    add("TOTAL OF AMINO ACIDS FOR SYNTHESIS")
    table([r for p in peptides for r in p.seq])

    maxlen = max(p.length for p in peptides)
    for w in range(math.ceil(maxlen / 10)):
        lo, hi = w * 10, (w + 1) * 10
        add(f"CANTIDAD DE AMINOACIDOS DESDE ACOPLE {lo + 1} AL {hi}")
        if windows_by == "posicion":
            res = [r for p in peptides for r in p.seq[lo:hi]]
            outs = sum(1 for p in peptides if lo < p.length <= hi)
        else:
            res = [r for p in peptides for r in list(reversed(p.seq))[lo:hi]]
            outs = sum(1 for p in peptides if lo < p.length + 1 <= hi)
        table(res, outs)
        add(); add(); add()
    for _ in range(4):
        add()

    def field_line(text, boxed, col=45, fields=""):
        t = text + (CHK if boxed else "")
        return _plain(t.ljust(col) + fields)

    for k in range(1, maxlen + 2):
        groups, summary = {}, []
        for p in peptides:
            rev = list(reversed(p.seq))
            if k <= p.length:
                res = rev[k - 1]
                groups.setdefault(res.label, (res.sortkey, []))[1].append(p.bag)
                summary.append((p.bag, res.label, k == p.length))
            elif k == p.length + 1:
                groups.setdefault("OUT", (("~", False), []))[1].append(p.bag)
                summary.append((p.bag, "OUT", False))
        L.append([PAGEBREAK])
        add(f"ACOPLE NUMERO {k}")
        for lab, (_, bags) in sorted(groups.items(), key=lambda kv: kv[1][0]):
            L.append([(f"        {lab + ':':<8}{','.join(map(str, bags))} =", ""),
                      (str(len(bags)), "b")])
        add("DESPROTECCION  FECHA ___/___/___            CHEQUEO   HECHO POR     REVISADO POR ")
        add()
        L.append(field_line(f"{deprotection} (2x10')", True, 43, "|____|___|______________ ______________"))
        add()
        L.append(field_line("    Lavado con DMF(3x1')", True, 43, "|__|__|__|______________ ______________"))
        add()
        L.append(field_line("    Lavado con IPA(1x1')", False, 43, "|________|______________ ______________"))
        add()
        L.append(field_line("Lavado con BPB 1%/DMF(1x2')", False, 43, "|________|______________ ______________"))
        add()
        L.append(field_line("    Lavado con DMF (2x1')", True, 43, "|____|___|______________ ______________"))
        add()
        L.append(field_line("    Lavado con DCM (1x1')", False, 43, "|________|______________ ______________"))
        add()
        add("Ciclo de Acople               FECHA     HORA          HECHO POR      REVISADO POR ")
        for nombre, boxed in (("simple", False), ("doble", True), ("triple", True)):
            texto, prop = ACTIVATORS[couplings[nombre].upper()]
            add(f"        {texto}")
            L.append(field_line(f"    {nombre.capitalize()} {prop}", boxed, 27,
                                "|___/___/___|__:__|__:__|______________ ______________"))
        add("Despues de cada ciclo y antes del test BPB lavar con DMF(2x1')" + CHK)
        for bag, lab, last in summary:
            L.append([(f"{bag}:{lab}", "u" if last else ""), ("   ", "")])
    return L


def build_reactor_program(peptide, reactor_id, name, mg, deprotection, couplings=None):
    couplings = {**DEFAULT_COUPLINGS, **(couplings or {})}
    L = []
    add = lambda s="": L.append(_plain(s))

    mw_str = f"{peptide.mw:.2f}" if peptide.mw is not None else "s/d"

    add(f"Nombre de Síntesis: {name}")
    add(f"Cantidad de Síntesis: {mg} mg.")
    add(f"Método de Desprotección: {deprotection} ")
    add(f"{'Largo':<12}{'M.W.':<11}{'ID Reactor':<10}")
    add(f"{peptide.length:<12}{mw_str:<11}{reactor_id:<10}")
    add(f"Secuencia: {peptide.text}")

    rev_seq = list(reversed(peptide.seq))
    coupling_steps = [r.label for r in rev_seq] + ["OUT"]

    for i, aa_label in enumerate(coupling_steps, start=1):
        if i > 1:
            L.append([PAGEBREAK])
        add(f"ACOPLE NUMERO {i}")
        L.append([(f"        {aa_label + ':':<9}= ", ""), (str(reactor_id), "b")])
        add("DESPROTECCION  FECHA ___/___/___         CHEQUEO   HECHO POR     REVISADO POR ")
        add(f"{deprotection} (2x10')       |____|___|______________ ______________")
        add("    Lavado con DMF(3x1')               |__|__|__|______________ ______________")
        add("    Ensayo ninhidrina                  |________|______________ ______________")
        add("    Lavado con DCM (1x1')              |________|______________ ______________")
        add("Ciclo de Acople               FECHA     HORA          HECHO POR      REVISADO POR ")
        
        for nombre in ("simple", "doble", "triple"):
            texto, prop = ACTIVATORS[couplings[nombre].upper()]
            add(f"        {texto}")
            add(f"    {nombre.capitalize()} {prop:<18}|___/___/___|__:__|__:__|______________ ______________")
            
        add("Despues de cada ciclo y antes del test BPB lavar con DMF(2x1')")
        add(f":{aa_label}    ")

    return L


def generate_docx_bytes(lines, title):
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.oxml.ns import qn

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin, sec.right_margin = Inches(0.79), Inches(0.39)
    sec.top_margin = sec.bottom_margin = Inches(0.39)
    st_doc = doc.styles["Normal"]
    st_doc.font.name = "Courier New"
    st_doc.element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    st_doc.font.size = Pt(10)
    st_doc.paragraph_format.space_after = Pt(0)
    st_doc.paragraph_format.space_before = Pt(0)
    st_doc.paragraph_format.line_spacing = 1.0
    doc.core_properties.title = title

    pending_break = False
    for ln in lines:
        if ln == [PAGEBREAK]:
            pending_break = True
            continue
        p = doc.add_paragraph()
        if pending_break:
            p.paragraph_format.page_break_before = True
            pending_break = False
        for text, style in ln:
            r = p.add_run(text)
            r.bold = style == "b"
            r.underline = style == "u"

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_pdf_bytes(lines, title):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase.pdfmetrics import stringWidth

    buffer = io.BytesIO()
    W, H = letter
    left, top, bottom, size, lead = 0.79 * 72, H - 0.39 * 72, 0.39 * 72, 10, 11.6
    c = canvas.Canvas(buffer, pagesize=letter)
    c.setTitle(title)
    y = top - size

    def newpage():
        nonlocal y
        c.showPage()
        y = top - size

    for ln in lines:
        if ln == [PAGEBREAK]:
            newpage()
            continue
        if y < bottom:
            newpage()
        x = left
        for text, style in ln:
            font = "Courier-Bold" if style == "b" else "Courier"
            c.setFont(font, size)
            c.drawString(x, y, text)
            w = stringWidth(text, font, size)
            if style == "u" and text.strip():
                c.line(x, y - 1.5, x + stringWidth(text.rstrip(), font, size), y - 1.5)
            x += w
        y -= lead
    c.save()
    buffer.seek(0)
    return buffer


def parse_dataframe(df):
    cols = {str(c).strip().lower(): c for c in df.columns}
    pick = lambda *names: next((cols[n] for n in names if n in cols), None)
    c_seq = pick("secuencia", "sequence", "seq")
    c_bag = pick("bolsa", "bag")
    c_fam = pick("familia", "family")
    c_pos = pick("pos", "pos.", "posicion", "posición")
    
    if c_seq is None:
        raise ValueError("El archivo cargado debe contener la columna 'Secuencia'")
    
    peps = []
    for i, row in df.iterrows():
        if str(row[c_seq]).strip() in ("", "nan"):
            continue
        clean = lambda v: "" if str(v) == "nan" else str(v).removesuffix(".0")
        bag = int(row[c_bag]) if c_bag and str(row[c_bag]) != "nan" else len(peps) + 1
        text = re.sub(r"\s+", "", str(row[c_seq]))
        peps.append(Peptide(bag, parse_sequence(text), text,
                            clean(row[c_fam]) if c_fam else "",
                            clean(row[c_pos]) if c_pos else ""))
    return peps


# --------------------------------------------------------------------------
# Componentes Visuales y Barra Lateral
# --------------------------------------------------------------------------
with st.sidebar:
    st.markdown("### 🧬 **SFS Studio**")
    st.caption("Generación de Programas de Sintesis.")
    st.divider()

    st.markdown("#### **Información del Sistema**")
    st.info("Plataforma para la generación de programas de uso en la Sintesis en fase solida Fmoc.")

    st.divider()
    st.markdown("Desarrollado en Streamlit y Python")
    st.caption("Javier Badilla.")

# Banner Principal
col_logo, col_header = st.columns([1, 6])
with col_logo:
    st.title("🧪")
with col_header:
    st.title("Generador de Programas de Síntesis")
    st.caption("Plataforma interactiva para la creación de programas de síntesis simultánea Tea Bag y en reactor.")

st.markdown("---")

tab1, tab2 = st.tabs(["📊 **Síntesis Simultánea (Bolsas)**", "⚗️ **Síntesis en Reactor**"])

# --------------------------------------------------------------------------
# TAB 1: Simultánea
# --------------------------------------------------------------------------
with tab1:
    st.subheader("Configuración")
    st.write("Cargue un archivo Excel/CSV con las columnas `Bolsa`, `Secuencia`, `Familia` y `Pos`.")

    with st.container(border=True):
        c1, c2, c3 = st.columns(3)
        with c1:
            s_nombre = st.text_input("ID / Nombre de Síntesis", value=datetime.date.today().strftime("S%m%d%Y"))
        with c2:
            s_mg = st.number_input("Masa resina por bolsa (mg)", value=40, step=5)
        with c3:
            s_desprot = st.text_input("Método de Desprotección", value=DEFAULT_DEPROTECTION)

    uploaded_file = st.file_uploader("Subir planilla de péptidos (.xlsx, .csv)", type=["xlsx", "csv"])

    if uploaded_file is not None:
        try:
            df = pd.read_csv(uploaded_file) if uploaded_file.name.endswith(".csv") else pd.read_excel(uploaded_file)
            peptides = parse_dataframe(df)

            st.success(f"Se cargaron **{len(peptides)} péptidos** correctamente.")
            
            with st.expander("👁️ Ver datos cargados"):
                st.dataframe(df, use_container_width=True)

            if st.button("🚀 Generar Documentos", type="primary", key="btn_sim"):
                with st.spinner("Procesando y generando archivos Word/PDF..."):
                    lines = build_simultaneous_program(peptides, s_nombre, s_mg, s_desprot, DEFAULT_COUPLINGS)
                    docx_buf = generate_docx_bytes(lines, s_nombre)
                    pdf_buf = generate_pdf_bytes(lines, s_nombre)

                st.markdown("### 📥 Archivos Listos para Descarga")
                d1, d2 = st.columns(2)
                with d1:
                    st.download_button(
                        label="📄 Descargar Programa Word (.docx)",
                        data=docx_buf,
                        file_name=f"{s_nombre}_simultaneo.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )
                with d2:
                    st.download_button(
                        label="📕 Descargar Programa PDF (.pdf)",
                        data=pdf_buf,
                        file_name=f"{s_nombre}_simultaneo.pdf",
                        mime="application/pdf",
                        use_container_width=True
                    )

        except Exception as e:
            st.error(f"Error procesando la entrada: {e}")

# --------------------------------------------------------------------------
# TAB 2: Reactor
# --------------------------------------------------------------------------
with tab2:
    st.subheader("Configuración de Reactor")
    
    with st.container(border=True):
        rc1, rc2 = st.columns(2)
        with rc1:
            r_nombre = st.text_input("Nombre de Síntesis", value=datetime.date.today().strftime("S%m%d%Y"), key="r_nom")
            r_secuencia = st.text_input("Secuencia (1 letra)", value="FFGDPDKDGTIDLKE")
        with rc2:
            r_reactor = st.text_input("ID Reactor", value="5")
            r_mg = st.number_input("Resina en Reactor (mg)", value=40, step=5, key="r_mg_val")

        r_desprot = st.text_input("Método de Desprotección ", value="PP 20% TritonX100 1%/DMF", key="r_desp")

    if st.button("🚀 Generar Programa Reactor", type="primary", key="btn_reac"):
        try:
            with st.spinner("Calculando peso molecular y secuencias..."):
                pep = Peptide(1, parse_sequence(r_secuencia), r_secuencia)
                lines = build_reactor_program(pep, r_reactor, r_nombre, r_mg, r_desprot, DEFAULT_COUPLINGS)

                docx_buf = generate_docx_bytes(lines, r_nombre)
                pdf_buf = generate_pdf_bytes(lines, r_nombre)

            st.success("¡Programa de reactor generado correctamente!")
            
            rd1, rd2 = st.columns(2)
            with rd1:
                st.download_button(
                    label="📄 Descargar en Word (.docx)",
                    data=docx_buf,
                    file_name=f"{r_nombre}_reactor.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )
            with rd2:
                st.download_button(
                    label="📕 Descargar en PDF (.pdf)",
                    data=pdf_buf,
                    file_name=f"{r_nombre}_reactor.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )

        except Exception as e:
            st.error(f"Error en los datos del reactor: {e}")